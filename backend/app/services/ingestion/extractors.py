"""Read source data only. Never launch Office, execute macros, or fetch external links."""

import csv
import json
import re
import warnings
import zipfile
from abc import ABC, abstractmethod
from itertools import islice
from pathlib import PurePosixPath
from typing import Protocol

from defusedxml.ElementTree import fromstring
from PIL import Image

from app.schemas.sources import ExtractionResult, Section

MAX_TEXT = 200_000
MAX_SECTIONS = 500
MAX_ROWS = 100
MAX_COLUMNS = 30


class ExtractionError(ValueError):
    pass


class VisionOCRAdapter(Protocol):
    def extract_text(self, path) -> str: ...


class Extractor(ABC):
    @abstractmethod
    def extract(self, path, title: str) -> ExtractionResult: ...


def clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(text)).strip()


def normalize(title, sections, metadata=None):
    kept, budget = [], MAX_TEXT
    metadata = dict(metadata or {})
    for section in sections:
        if len(kept) >= MAX_SECTIONS or budget <= 0:
            metadata["truncated"] = True
            break
        text = clean(section.text)
        heading = clean(section.title)[:500]
        available = max(0, budget - len(heading) - 4)
        if len(text) > available:
            metadata["truncated"] = True
        text = text[:available]
        kept.append(
            Section(title=heading, text=text, reference=section.reference, kind=section.kind)
        )
        budget -= len(heading) + len(text) + 4
    combined = "\n\n".join(
        "\n".join(filter(None, [section.title, section.text])) for section in kept
    )[:MAX_TEXT]
    return ExtractionResult(
        title=clean(title)[:500], text=combined, sections=kept, metadata=metadata
    )


def preflight_office(path, extension):
    required = {
        "docx": "word/document.xml",
        "pptx": "ppt/presentation.xml",
        "xlsx": "xl/workbook.xml",
    }
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 1500 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
            raise ExtractionError("Archive expansion limit exceeded")
        names = {item.filename for item in entries}
        if required[extension] not in names or "[Content_Types].xml" not in names:
            raise ExtractionError("Invalid Office package")
        if len(names) != len(entries):
            raise ExtractionError("Duplicate archive members")
        for entry in entries:
            name = entry.filename.replace("\\", "/")
            part = PurePosixPath(name)
            if part.is_absolute() or ".." in part.parts or ":" in name:
                raise ExtractionError("Unsafe archive member")
            if ((entry.external_attr >> 16) & 0o170000) == 0o120000 or entry.flag_bits & 1:
                raise ExtractionError("Encrypted or linked archive member")
            if entry.file_size > 10 * 1024 * 1024 or (
                entry.file_size > max(entry.compress_size, 1) * 200
            ):
                raise ExtractionError("Archive expansion limit exceeded")
            lower = name.lower()
            if any(marker in lower for marker in ("vbaproject", "/activex/", "/embeddings/")):
                raise ExtractionError("Active or embedded Office content is not supported")
            # CRC validation and bounds apply to all members; never extract to disk.
            data = archive.read(entry)
            if lower.endswith((".xml", ".rels")):
                fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)


class PDFExtractor(Extractor):
    def extract(self, path, title):
        from pypdf import PdfReader, overwrite_configuration

        overwrite_configuration(zlib_maximum_output_length=8 * 1024 * 1024)
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ExtractionError("Password-protected PDF is not supported")
        total = len(reader.pages)
        sections = []
        for index, page in enumerate(islice(reader.pages, 200), start=1):
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 8 * 1024 * 1024:
                raise ExtractionError("PDF page content limit exceeded")
            text = page.extract_text() or ""
            sections.append(
                Section(
                    title="",
                    text=text[:MAX_TEXT],
                    reference={"page": index},
                    kind="page",
                )
            )
            if sum(len(section.text) for section in sections) >= MAX_TEXT:
                break
        metadata = {
            "pages": total,
            "extracted_pages": len(sections),
            "truncated": total > len(sections),
        }
        result = normalize(title, sections, metadata)
        if not result.text:
            result.metadata["ocr"] = "unavailable"
            result.metadata["warning"] = "No text layer was found; OCR is not configured."
        return result


class DocxExtractor(Extractor):
    def extract(self, path, title):
        from docx import Document
        from docx.table import Table

        document = Document(path)
        sections, lines, heading = [], [], ""
        truncated = False
        for index, block in enumerate(document.iter_inner_content()):
            if index >= 5000 or sum(len(line) for line in lines) > MAX_TEXT:
                truncated = True
                break
            if isinstance(block, Table):
                lines.append(
                    "\n".join(
                        "\t".join(cell.text[:500] for cell in row.cells[:MAX_COLUMNS])
                        for row in block.rows[:MAX_ROWS]
                    )
                )
                continue
            style = block.style.name if block.style else ""
            if style.startswith("Heading") or style == "Title":
                if heading or lines:
                    sections.append(
                        Section(
                            title=heading[:500],
                            text="\n".join(lines)[:MAX_TEXT],
                            reference={"section": len(sections) + 1},
                            kind="heading",
                        )
                    )
                heading, lines = block.text, []
            else:
                bullet = style.startswith("List") or bool(block._p.xpath("./w:pPr/w:numPr"))
                lines.append(("- " if bullet else "") + block.text)
        sections.append(
            Section(
                title=heading[:500],
                text="\n".join(lines)[:MAX_TEXT],
                reference={"section": len(sections) + 1},
                kind="heading" if heading else "section",
            )
        )
        return normalize(
            document.core_properties.title or title, sections, {"truncated": truncated}
        )


class PptxExtractor(Extractor):
    def extract(self, path, title):
        from pptx import Presentation

        presentation = Presentation(path)
        sections = []
        for index, slide in enumerate(islice(presentation.slides, 200), start=1):
            heading = slide.shapes.title.text if slide.shapes.title else f"Slide {index}"
            lines = []

            def visit(shapes, depth=0, *, slide=slide, lines=lines):
                if depth > 5:
                    return
                for shape in islice(shapes, 1000):
                    if shape == slide.shapes.title:
                        continue
                    if shape.has_text_frame:
                        lines.append(shape.text[:20_000])
                    if shape.has_table:
                        lines.extend(
                            "\t".join(cell.text[:500] for cell in islice(row.cells, MAX_COLUMNS))
                            for row in islice(shape.table.rows, MAX_ROWS)
                        )
                    if hasattr(shape, "shapes"):
                        visit(shape.shapes, depth + 1)

            visit(slide.shapes)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    lines.append("Speaker notes: " + notes[:20_000])
            sections.append(
                Section(
                    title=heading[:500],
                    text="\n".join(lines)[:MAX_TEXT],
                    reference={"slide": index},
                    kind="slide",
                )
            )
        return normalize(
            title,
            sections,
            {
                "slides": len(presentation.slides),
                "truncated": len(presentation.slides) > 200,
            },
        )


def table_section(rows, name):
    rows = list(islice(rows, MAX_ROWS + 2))
    has_more = len(rows) > MAX_ROWS + 1 or any(len(row) > MAX_COLUMNS for row in rows)
    rows = rows[: MAX_ROWS + 1]
    lines = [
        "\t".join(clean(value if value is not None else "")[:500] for value in row[:MAX_COLUMNS])
        for row in rows
    ]
    return Section(
        title=name[:500],
        text="\n".join(lines)[:MAX_TEXT],
        kind="table",
        reference={"sheet": name[:200], "row_start": 1, "row_end": len(rows)},
    ), has_more


class SpreadsheetExtractor(Extractor):
    def extract(self, path, title):
        if path.suffix.lower() == ".csv":
            csv.field_size_limit(20_000)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                section, truncated = table_section(csv.reader(handle, strict=True), "CSV")
            return normalize(
                title,
                [section],
                {
                    "sheets": 1,
                    "truncated": truncated,
                    "row_limit": MAX_ROWS,
                    "column_limit": MAX_COLUMNS,
                    "formulas_executed": False,
                },
            )
        from openpyxl import load_workbook

        workbook = load_workbook(
            path, read_only=True, data_only=True, keep_vba=False, keep_links=False
        )
        try:
            sections, truncated = [], len(workbook.worksheets) > 20
            for sheet in workbook.worksheets[:20]:
                section, more = table_section(
                    sheet.iter_rows(
                        max_row=min(sheet.max_row or MAX_ROWS + 2, MAX_ROWS + 2),
                        max_col=min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS),
                        values_only=True,
                    ),
                    sheet.title,
                )
                sections.append(section)
                truncated = truncated or more or (sheet.max_column or 0) > MAX_COLUMNS
            return normalize(
                title,
                sections,
                {
                    "sheets": len(workbook.worksheets),
                    "truncated": truncated,
                    "row_limit": MAX_ROWS,
                    "column_limit": MAX_COLUMNS,
                    "formula_policy": "Cached values only; not evaluated",
                },
            )
        finally:
            workbook.close()


class TextExtractor(Extractor):
    def extract(self, path, title):
        with path.open(encoding="utf-8-sig") as handle:
            text = handle.read(MAX_TEXT + 1)
        if "\x00" in text:
            raise ExtractionError("Binary data is not a text source")
        sections, heading, lines = [], "", []
        for line in text[:MAX_TEXT].splitlines():
            if path.suffix.lower() == ".md" and re.match(r"^#{1,6}\s+", line):
                if heading or lines:
                    sections.append(Section(title=heading, text="\n".join(lines), kind="heading"))
                heading, lines = re.sub(r"^#{1,6}\s+", "", line)[:500], []
            else:
                lines.append(line)
        sections.append(
            Section(title=heading, text="\n".join(lines), kind="heading" if heading else "text")
        )
        for index, section in enumerate(sections):
            section.reference = {"section": index + 1}
        return normalize(title, sections, {"encoding": "utf-8", "truncated": len(text) > MAX_TEXT})


class JsonExtractor(Extractor):
    def extract(self, path, title):
        def invalid_constant(value):
            raise ExtractionError("Non-finite JSON number")

        with path.open(encoding="utf-8-sig") as handle:
            data = json.load(handle, parse_constant=invalid_constant)
        lines, remaining = [], MAX_TEXT
        truncated = False

        def visit(value, key="$", depth=0):
            nonlocal remaining, truncated
            if depth > 40:
                raise ExtractionError("JSON nesting limit exceeded")
            if remaining <= 0 or len(lines) >= 5000:
                truncated = True
                return
            if isinstance(value, dict):
                for index, (name, child) in enumerate(value.items()):
                    if index >= 1000 or remaining <= 0:
                        truncated = True
                        break
                    visit(child, key + "." + str(name)[:200], depth + 1)
            elif isinstance(value, list):
                for index, child in enumerate(value[:1000]):
                    visit(child, f"{key}[{index}]", depth + 1)
                truncated = truncated or len(value) > 1000
            else:
                line = f"{key}: {json.dumps(value, ensure_ascii=False)}"
                lines.append(line[:remaining])
                remaining -= len(line) + 1

        visit(data)
        return normalize(
            title,
            [
                Section(
                    title="",
                    text="\n".join(lines)[:MAX_TEXT],
                    reference={"json_path": "$"},
                    kind="json",
                )
            ],
            {"truncated": truncated},
        )


class ImageExtractor(Extractor):
    def __init__(self, ocr: VisionOCRAdapter | None = None):
        self.ocr = ocr

    def extract(self, path, title):
        Image.MAX_IMAGE_PIXELS = 25_000_000
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as image:
                expected = {"png": "PNG", "jpg": "JPEG", "jpeg": "JPEG", "webp": "WEBP"}
                if image.format != expected[path.suffix[1:].lower()]:
                    raise ExtractionError("Image encoding does not match extension")
                metadata = {
                    "width": image.width,
                    "height": image.height,
                    "format": image.format,
                    "mode": image.mode,
                    "ocr": "configured" if self.ocr else "unavailable",
                }
                image.verify()
            with Image.open(path) as image:
                image.load()
        text = self.ocr.extract_text(path) if self.ocr else ""
        return normalize(
            title,
            [Section(text=text, reference={"image": 1}, kind="image")] if text else [],
            metadata,
        )


EXTRACTORS = {
    "pdf": PDFExtractor,
    "docx": DocxExtractor,
    "pptx": PptxExtractor,
    "xlsx": SpreadsheetExtractor,
    "csv": SpreadsheetExtractor,
    "txt": TextExtractor,
    "md": TextExtractor,
    "json": JsonExtractor,
    "png": ImageExtractor,
    "jpg": ImageExtractor,
    "jpeg": ImageExtractor,
    "webp": ImageExtractor,
}


def extract_document(path, extension, title):
    if extension not in EXTRACTORS:
        raise ExtractionError("Unsupported source format")
    if extension in {"docx", "pptx", "xlsx"}:
        preflight_office(path, extension)
    return EXTRACTORS[extension]().extract(path, title)
