from io import BytesIO

from pptx import Presentation


async def test_export_uses_unsaved_content_without_mutating_deck(client):
    created = await client.post("/presentations", json={"title": "Original", "markdown": "# Old"})
    deck_id = created.json()["id"]
    exported = await client.post(
        f"/presentations/{deck_id}/export/pptx",
        json={"title": "Current", "markdown": "# One\n\nBody\n\n---\n\n# Two"},
    )
    assert exported.status_code == 200
    assert "presentationml.presentation" in exported.headers["content-type"]
    deck = Presentation(BytesIO(exported.content))
    assert len(deck.slides) == 2
    assert deck.core_properties.title == "Current"
    assert deck.slides[0].shapes[0].text == "One"
    assert "Body" in deck.slides[0].notes_slide.notes_text_frame.text
    persisted = await client.get(f"/presentations/{deck_id}")
    assert persisted.json()["markdown"] == "# Old"
