"""Small real files shared by ingestion tests; no downloaded or secret fixtures."""

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def source_bytes(extension):
    stream = BytesIO()
    if extension == "pdf":
        writer = PdfWriter()
        page = writer.add_blank_page(612, 792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 700 Td (Revenue grew 12 percent in 2025.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        writer.write(stream)
    elif extension == "docx":
        doc = Document()
        doc.add_heading("Summary", level=1)
        doc.add_paragraph("Revenue grew 12 percent in 2025.")
        doc.add_paragraph("Supported finding", style="List Bullet")
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text, cells[1].text = "Region", "North"
        doc.save(stream)
    elif extension == "pptx":
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Summary"
        slide.placeholders[1].text = "Revenue grew 12 percent in 2025."
        slide.notes_slide.notes_text_frame.text = "Verified by finance"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(5), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "Region", "North"
        deck.save(stream)
    elif extension == "xlsx":
        book = Workbook()
        book.active.title = "Revenue"
        book.active.append(["Year", "Growth"])
        book.active.append([2025, 12])
        book.save(stream)
        book.close()
    elif extension in {"png", "jpg", "jpeg", "webp"}:
        Image.new("RGB", (16, 12), "white").save(
            stream, format={"jpg": "JPEG", "jpeg": "JPEG"}.get(extension, extension.upper())
        )
    else:
        return {
            "csv": b"Year,Growth\n2025,12\n",
            "txt": b"Revenue grew 12 percent in 2025.",
            "md": b"# Summary\nRevenue grew 12 percent in 2025.",
            "json": b'{"revenue": {"growth": 12, "year": 2025}}',
        }[extension]
    return stream.getvalue()
