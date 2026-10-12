from io import BytesIO
from uuid import uuid4

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.pptx_export import export_pptx


def test_pptx_preserves_palette_stored_image_notes_and_aspect(tmp_path):
    asset_id = str(uuid4())
    path = tmp_path / "asset.png"
    Image.new("RGB", (800, 400), "#226688").save(path)
    markdown = (
        "# A clear story\n\n- First point\n\n"
        f"![Research](/api/v1/media/{asset_id}/file)\n\n"
        "<!-- speaker notes\nKeep this context\n-->"
    )
    payload = export_pptx(
        "Visual deck", markdown,
        {"background": "#112233", "text": "#ffffff", "accent": "#22aacc"},
        {asset_id: path},
    )
    deck = Presentation(BytesIO(payload))
    slide = deck.slides[0]
    assert str(slide.background.fill.fore_color.rgb) == "112233"
    assert str(slide.shapes[0].text_frame.paragraphs[0].font.color.rgb) == "FFFFFF"
    picture = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert abs(picture.width / picture.height - 2) < 0.01
    assert "Keep this context" in slide.notes_slide.notes_text_frame.text
    assert deck.slide_width / deck.slide_height > 1.77
